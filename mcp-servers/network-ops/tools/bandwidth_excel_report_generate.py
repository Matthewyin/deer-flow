import uuid
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional

from fastmcp import FastMCP
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import get_config
from db.bandwidth_lines_client import BandwidthLinesClient


_client: Optional[BandwidthLinesClient] = None
_PROJECT_ROOT = Path(__file__).resolve().parents[3]
_VIRTUAL_OUTPUT_PREFIX = "/mnt/user-data/outputs/.mcp/network-ops"
_DEFAULT_LINE_NOS = (5, 6, 151, 152, 153, 154, 159, 160, 161, 162, 201, 202, 203)
_GROUP_ORDER = {
    "TLS终端专线": 1,
    "北单售票专线": 2,
    "体彩APP专线": 3,
    "西五环互联网B区线路": 4,
}


def _get_client() -> BandwidthLinesClient:
    global _client
    if _client is None:
        config = get_config()
        _client = BandwidthLinesClient(config.sqlite.db_path)
    return _client


def _resolve_path(path_str: str) -> Path:
    path = Path(path_str)
    if path.is_absolute():
        return path
    return _PROJECT_ROOT / path


def _safe_output_filename(filename: str) -> str:
    name = Path((filename or "").strip()).name
    if not name:
        name = "带宽峰值均值日报表.xlsx"
    if not name.endswith(".xlsx"):
        name += ".xlsx"
    return name


def _resolve_period(days: int, start_date: str, end_date: str, available_dates: list[str]):
    if not available_dates:
        return "", "", []

    end = end_date or max(available_dates)
    if start_date:
        start = start_date
    else:
        safe_days = max(1, min(days or 7, 31))
        end_dt = datetime.strptime(end, "%Y-%m-%d")
        start = (end_dt - timedelta(days=safe_days - 1)).strftime("%Y-%m-%d")

    dates = [day for day in available_dates if start <= day <= end]
    return start, end, dates


def _to_float(value) -> float:
    if value is None:
        return 0.0
    return round(float(value), 2)


def _avg_util(avg_mbps, bandwidth_mbps) -> float:
    bandwidth = _to_float(bandwidth_mbps)
    if bandwidth <= 0:
        return 0.0
    return round(_to_float(avg_mbps) * 100 / bandwidth, 2)


def _line_sort_key(row: dict):
    return (
        _GROUP_ORDER.get(row.get("line_group") or "", 99),
        row.get("line_no") or 0,
    )


def _build_rows(rows: list[dict], day: str) -> list[list]:
    day_rows = [row for row in rows if row.get("report_date") == day]
    result = []
    for row in sorted(day_rows, key=_line_sort_key):
        bandwidth = row.get("bandwidth_mbps") or 0
        result.append(
            [
                row.get("line_group") or "",
                row.get("line_no") or "",
                row.get("carrier") or "",
                bandwidth,
                round(float(bandwidth) * 0.8, 2) if bandwidth else 0,
                _to_float(row.get("in_peak_mbps")),
                _to_float(row.get("in_peak_util_pct")),
                row.get("in_peak_time") or "",
                _to_float(row.get("out_peak_mbps")),
                _to_float(row.get("out_peak_util_pct")),
                row.get("out_peak_time") or "",
                _to_float(row.get("in_avg_mbps")),
                _avg_util(row.get("in_avg_mbps"), bandwidth),
                "-",
                _to_float(row.get("out_avg_mbps")),
                _avg_util(row.get("out_avg_mbps"), bandwidth),
                "-",
            ]
        )
    return result


def _write_workbook(rows: list[dict], dates: list[str], output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    headers = [
        "线路类型",
        "线路",
        "运营商",
        "带宽(M)",
        "阈值(M, 带宽*80%)",
        "入向峰值(Mbps)",
        "入向峰值利用率(%)",
        "入向峰值时间点",
        "出向峰值(Mbps)",
        "出向峰值利用率(%)",
        "出向峰值时间点",
        "入向均值(Mbps)",
        "入向均值利用率(%)",
        "入向均值时间点",
        "出向均值(Mbps)",
        "出向均值利用率(%)",
        "出向均值时间点",
    ]

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    group_fill = PatternFill("solid", fgColor="EAF2F8")

    for day in dates:
        ws = wb.create_sheet(day)
        ws.append(headers)
        for row in _build_rows(rows, day):
            ws.append(row)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")
            row[0].fill = group_fill

        for col_idx in range(1, len(headers) + 1):
            column = get_column_letter(col_idx)
            max_len = max(
                len(str(ws.cell(row=row_idx, column=col_idx).value or ""))
                for row_idx in range(1, ws.max_row + 1)
            )
            ws.column_dimensions[column].width = min(max(max_len + 2, 12), 24)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb.save(output_path)


def generate_bandwidth_excel_report(
    days: int = 7,
    start_date: str = "",
    end_date: str = "",
    output_filename: str = "",
) -> dict:
    client = _get_client()
    available_dates = client.get_available_dates()
    start, end, period_dates = _resolve_period(days, start_date, end_date, available_dates)
    if not period_dates:
        return {
            "ok": False,
            "error": "数据库中没有指定周期的线路带宽数据",
            "period": {"start": start, "end": end},
        }

    rows = [
        row
        for row in client.query_records(start, end)
        if row.get("line_no") in _DEFAULT_LINE_NOS
    ]
    expected = len(period_dates) * len(_DEFAULT_LINE_NOS)
    if len(rows) != expected:
        return {
            "ok": False,
            "error": "目标线路数据不完整，已停止生成 Excel",
            "period": {"start": start, "end": end},
            "expected_records": expected,
            "actual_records": len(rows),
            "line_nos": list(_DEFAULT_LINE_NOS),
        }

    cfg = get_config()
    filename = _safe_output_filename(output_filename)
    artifact_id = uuid.uuid4().hex
    output_dir = _resolve_path(cfg.bandwidth_report_output_dir) / artifact_id
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / filename
    virtual_path = f"{_VIRTUAL_OUTPUT_PREFIX}/{artifact_id}/{filename}"

    _write_workbook(rows, period_dates, output_path)
    return {
        "ok": True,
        "output_filename": filename,
        "artifact_path": str(output_path),
        "present_filepaths": [virtual_path],
        "period": {"start": start, "end": end},
        "date_count": len(period_dates),
        "sheet_count": len(period_dates),
        "line_count": len(_DEFAULT_LINE_NOS),
        "record_count": len(rows),
        "notes": ["源数据没有均值时间点字段，Excel 中均值时间点以 '-' 展示。"],
    }


def register(mcp: FastMCP):
    @mcp.tool()
    def bandwidth_excel_report_generate(
        days: int = 7,
        start_date: str = "",
        end_date: str = "",
        output_filename: str = "",
    ) -> dict:
        """生成默认 13 条专线近 N 天带宽峰值、峰值利用率、均值和均值利用率 Excel 表。

        Args:
            days: 统计天数，默认 7。仅在未提供 start_date 时生效。
            start_date: 开始日期（YYYY-MM-DD），空则根据 days 自动计算。
            end_date: 结束日期（YYYY-MM-DD），空则取最新可用日期。
            output_filename: 建议保存给用户的 xlsx 文件名。

        Returns:
            dict: 成功时返回 artifact_path、present_filepaths、period、sheet_count。
        """
        return generate_bandwidth_excel_report(
            days=days,
            start_date=start_date,
            end_date=end_date,
            output_filename=output_filename,
        )
