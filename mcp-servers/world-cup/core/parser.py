from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.normalizer import (
    infer_unit,
    infer_value_role,
    metric_key,
    normalize_date,
    normalize_metric_name,
    normalize_number,
    sheet_slug,
)

SUPPORTED_SHEETS = [
    "每小时数据—TPS、五书",
    "日统计数据—TPS、五书",
    "周统计数据—TPS、五书",
    "两网三端时段",
    "两网三端时段(日统计）",
]


def _value(ws, row: int, col: int) -> Any:
    return ws.cell(row=row, column=col).value


def _merged_parent_value(ws, row: int, col: int) -> Any:
    cell = ws.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value
    for merged in ws.merged_cells.ranges:
        if cell.coordinate in merged:
            return ws.cell(row=merged.min_row, column=merged.min_col).value
    return None


def _metric_group(ws, col: int) -> str:
    value = _merged_parent_value(ws, 1, col)
    return str(value).strip() if value else ""


def _record(
    *,
    sheet_name: str,
    granularity: str,
    row_no: int,
    source_column: str,
    period_key: str,
    date: str | None,
    week_label: str | None,
    hour_range: str,
    metric_group: str,
    header: str,
    value: Any,
) -> dict:
    role = infer_value_role(header)
    name = normalize_metric_name(header)
    numeric, text = normalize_number(value)
    if role == "compare_pct" and numeric is not None and isinstance(value, (int, float)) and abs(numeric) <= 1:
        numeric = numeric * 100
    return {
        "sheet_name": sheet_name,
        "sheet_slug": sheet_slug(sheet_name),
        "granularity": granularity,
        "period_key": period_key,
        "date": date,
        "week_label": week_label,
        "hour_range": hour_range,
        "metric_group": metric_group,
        "metric_key": metric_key(name),
        "metric_name": name,
        "value_role": role,
        "value_numeric": numeric,
        "value_text": text,
        "unit": infer_unit(name, role),
        "source_column": source_column,
        "row_no": row_no,
    }


def _parse_hourly_sheet(ws) -> list[dict]:
    records = []
    header_row = 1
    data_start_row = 2
    first_header = str(_value(ws, 1, 1) or "").strip()
    second_header = str(_value(ws, 1, 2) or "").strip()
    second_row_first = str(_value(ws, 2, 1) or "").strip()
    if (
        first_header != "日期"
        or second_header not in ("时间", "时段")
        or second_row_first == "日期"
    ):
        header_row = 2
        data_start_row = 3

    for row in range(data_start_row, ws.max_row + 1):
        date = normalize_date(_value(ws, row, 1))
        hour_range = str(_value(ws, row, 2) or "").strip()
        if not date and not hour_range:
            continue
        period_key = f"{date or ''} {hour_range}".strip()
        for col in range(3, ws.max_column + 1):
            header = _value(ws, header_row, col)
            if not header:
                continue
            records.append(
                _record(
                    sheet_name=ws.title,
                    granularity="hourly",
                    row_no=row,
                    source_column=get_column_letter(col),
                    period_key=period_key,
                    date=date,
                    week_label=None,
                    hour_range=hour_range,
                    metric_group=_metric_group(ws, col),
                    header=str(header),
                    value=_value(ws, row, col),
                )
            )
    return records


def _parse_flat_sheet(ws, granularity: str) -> list[dict]:
    records = []
    for row in range(2, ws.max_row + 1):
        first = _value(ws, row, 1)
        if first is None or first == "":
            continue
        date = normalize_date(first) if granularity == "daily" else None
        week_label = str(first).strip() if granularity == "weekly" else None
        period_key = date or week_label or str(first).strip()
        for col in range(2, ws.max_column + 1):
            header = _value(ws, 1, col)
            if not header:
                continue
            records.append(
                _record(
                    sheet_name=ws.title,
                    granularity=granularity,
                    row_no=row,
                    source_column=get_column_letter(col),
                    period_key=period_key,
                    date=date,
                    week_label=week_label,
                    hour_range="",
                    metric_group="",
                    header=str(header),
                    value=_value(ws, row, col),
                )
            )
    return records


def _parse_channel_hourly_sheet(ws) -> list[dict]:
    records = []
    for row in range(2, ws.max_row + 1):
        date = normalize_date(_value(ws, row, 1))
        hour_range = str(_value(ws, row, 2) or "").strip()
        if not date and not hour_range:
            continue
        period_key = f"{date or ''} {hour_range}".strip()
        for col in range(3, ws.max_column + 1):
            header = _value(ws, 1, col)
            if not header:
                continue
            records.append(
                _record(
                    sheet_name=ws.title,
                    granularity="hourly",
                    row_no=row,
                    source_column=get_column_letter(col),
                    period_key=period_key,
                    date=date,
                    week_label=None,
                    hour_range=hour_range,
                    metric_group="两网三端",
                    header=str(header),
                    value=_value(ws, row, col),
                )
            )
    return records


def parse_workbook(path: str | Path, upload_id: str, source_filename: str) -> dict:
    wb = load_workbook(path, read_only=False, data_only=True)
    records: list[dict] = []
    errors: list[dict] = []

    matched_sheets = [sheet_name for sheet_name in SUPPORTED_SHEETS if sheet_name in wb.sheetnames]
    if not matched_sheets:
        errors.append(
            {
                "upload_id": upload_id,
                "sheet_name": "",
                "row_no": None,
                "field_name": "sheet",
                "error_message": "未找到支持的工作表",
            }
        )

    for sheet_name in matched_sheets:
        ws = wb[sheet_name]
        try:
            if sheet_name == "每小时数据—TPS、五书":
                records.extend(_parse_hourly_sheet(ws))
            elif sheet_name == "日统计数据—TPS、五书":
                records.extend(_parse_flat_sheet(ws, "daily"))
            elif sheet_name == "周统计数据—TPS、五书":
                records.extend(_parse_flat_sheet(ws, "weekly"))
            elif sheet_name == "两网三端时段":
                records.extend(_parse_channel_hourly_sheet(ws))
            elif sheet_name == "两网三端时段(日统计）":
                records.extend(_parse_flat_sheet(ws, "daily"))
        except Exception as e:
            errors.append(
                {
                    "upload_id": upload_id,
                    "sheet_name": sheet_name,
                    "row_no": None,
                    "field_name": "sheet",
                    "error_message": str(e),
                }
            )

    for record in records:
        record["upload_id"] = upload_id

    return {
        "upload_id": upload_id,
        "source_filename": source_filename,
        "sheet_count": len(matched_sheets),
        "record_count": len(records),
        "records": records,
        "errors": errors,
    }
