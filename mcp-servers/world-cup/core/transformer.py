import shutil
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


HOURLY_SHEET = "每小时数据—TPS、五书"
CURRENT_HOURLY_HEADERS = [
    "日期",
    "时间",
    "平台TPS",
    "竞彩TPS",
    "传足TPS",
    "北单TPS",
    "超额申请",
    "高额预约",
    "大额预约",
    "竞彩票数",
    "竞彩销量",
    "传足票数",
    "传足销量",
    "北单票数",
    "北单销量",
]

OLD_TODAY_HEADER_BY_CURRENT = {
    "平台TPS": "平台TPS（今日）",
    "竞彩TPS": "竞彩TPS（今日）",
    "传足TPS": "传足TPS（今日）",
    "北单TPS": "北单TPS（今日）",
    "超额申请": "超额申请（今日）",
    "高额预约": "高额预约（今日）",
    "大额预约": "大额预约（今日）",
    "竞彩票数": "竞彩票数（今日）",
    "竞彩销量": "竞彩销量（今日）",
    "传足票数": "传足票数（今日）",
    "传足销量": "传足销量（今日）",
    "北单票数": "北单票数（今日）",
    "北单销量": "北单销量（今日）",
}


def _value(ws, row: int, col: int) -> Any:
    return ws.cell(row=row, column=col).value


def _bounds(ws) -> tuple[int, int]:
    # 部分旧表的 worksheet dimension 写成 A1，只能以实际加载到内存的单元格为准。
    cells = getattr(ws, "_cells", {})
    if cells:
        return max(row for row, _ in cells), max(col for _, col in cells)
    return ws.max_row, ws.max_column


def _is_old_hourly_sheet(ws) -> bool:
    _, max_col = _bounds(ws)
    second_row_headers = {str(_value(ws, 2, col) or "").strip() for col in range(1, max_col + 1)}
    return any(header.endswith("（今日）") for header in second_row_headers)


def _copy_sheet_values(source_ws, target_ws) -> None:
    max_row, max_col = _bounds(source_ws)
    for row in range(1, max_row + 1):
        target_ws.append([_value(source_ws, row, col) for col in range(1, max_col + 1)])


def _normalize_old_hourly_sheet(source_ws, target_ws) -> None:
    max_row, max_col = _bounds(source_ws)
    source_headers = {
        str(_value(source_ws, 2, col) or "").strip(): col
        for col in range(1, max_col + 1)
        if _value(source_ws, 2, col) is not None
    }

    target_ws.append(CURRENT_HOURLY_HEADERS)
    for row in range(3, max_row + 1):
        date = _value(source_ws, row, 1)
        hour_range = _value(source_ws, row, 2)
        if (date is None or date == "") and (hour_range is None or hour_range == ""):
            continue

        values = [date, hour_range]
        for current_header in CURRENT_HOURLY_HEADERS[2:]:
            source_col = source_headers.get(OLD_TODAY_HEADER_BY_CURRENT[current_header])
            values.append(_value(source_ws, row, source_col) if source_col else "")
        target_ws.append(values)


def normalize_workbook_for_import(source_path: str | Path, target_path: str | Path) -> dict:
    source = Path(source_path)
    target = Path(target_path)
    target.parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(source, read_only=False, data_only=True)
    if HOURLY_SHEET not in wb.sheetnames or not _is_old_hourly_sheet(wb[HOURLY_SHEET]):
        shutil.copyfile(source, target)
        return {"normalized": False, "path": str(target)}

    normalized = Workbook()
    normalized.remove(normalized.active)

    for sheet_name in wb.sheetnames:
        target_ws = normalized.create_sheet(sheet_name)
        source_ws = wb[sheet_name]
        if sheet_name == HOURLY_SHEET:
            _normalize_old_hourly_sheet(source_ws, target_ws)
        else:
            _copy_sheet_values(source_ws, target_ws)

    normalized.save(target)
    return {"normalized": True, "path": str(target)}
