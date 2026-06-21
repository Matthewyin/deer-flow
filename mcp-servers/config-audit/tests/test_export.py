import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.export import export_report  # noqa: E402
from core.model import (  # noqa: E402
    AddressObject,
    CompareResult,
    ConfigFinding,
    DeviceProfile,
    Interface,
    ManagementAccess,
    NormalizedConfig,
    PolicyRule,
    ServiceObject,
    Zone,
)


def test_export_report_writes_excel_markdown_and_present_paths(tmp_path):
    config = _config()
    compare = CompareResult(
        template_id="firewall.dmz.border_firewall",
        findings=[
            ConfigFinding(
                finding_id="F-0001",
                severity="high",
                module="policy_rules",
                finding_type="risky",
                expected="禁止 any 到 any 的放行策略",
                actual="策略 allow_any 允许 any 到 any",
                evidence="security-policy allow_any",
                llm_analysis="该策略暴露面过大。",
                recommendation="收敛源、目的和服务范围，避免 any 到 any 放行",
            )
        ],
    )

    paths = export_report(
        config,
        compare,
        tmp_path,
        agent_analysis="本次分析由 Agent 汇总，建议优先处理高危策略。",
    )

    excel_path = Path(paths.excel_path)
    markdown_path = Path(paths.markdown_path)
    output_root = tmp_path / "mcp-outputs"

    assert excel_path.exists()
    assert markdown_path.exists()
    excel_relative = excel_path.relative_to(output_root).as_posix()
    markdown_relative = markdown_path.relative_to(output_root).as_posix()
    assert excel_relative == f"config-audit/{excel_path.parent.name}/配置审查事实表.xlsx"
    assert markdown_relative == f"config-audit/{excel_path.parent.name}/配置审查报告.md"
    assert paths.present_filepaths == [
        f"/mnt/user-data/outputs/.mcp/{excel_relative}",
        f"/mnt/user-data/outputs/.mcp/{markdown_relative}",
    ]

    workbook = load_workbook(excel_path)
    assert {
        "device_profile",
        "zones",
        "interfaces",
        "address_objects",
        "service_objects",
        "policy_rules",
        "nat_rules",
        "routes",
        "management_access",
        "logging",
        "unparsed_blocks",
        "findings",
    }.issubset(set(workbook.sheetnames))

    findings = workbook["findings"]
    assert [cell.value for cell in findings[1]] == [
        "finding_id",
        "severity",
        "module",
        "finding_type",
        "expected",
        "actual",
        "evidence",
        "llm_analysis",
        "recommendation",
    ]
    assert findings["A2"].value == "F-0001"
    assert findings["I2"].value == "收敛源、目的和服务范围，避免 any 到 any 放行"

    nat_rules = workbook["nat_rules"]
    assert nat_rules.max_row == 2
    assert [cell.value for cell in nat_rules[1]] == [
        "name",
        "nat_type",
        "source",
        "destination",
        "translated",
        "raw",
    ]
    assert nat_rules["A2"].value == "无数据"

    markdown = markdown_path.read_text(encoding="utf-8")
    assert "# 配置审查报告" in markdown
    assert "本次分析由 Agent 汇总，建议优先处理高危策略。" in markdown
    assert "收敛源、目的和服务范围，避免 any 到 any 放行" in markdown


def test_export_report_normalizes_output_root_variants(tmp_path):
    config = _config()
    compare = CompareResult(template_id="firewall.dmz.border_firewall")

    plain_paths = export_report(config, compare, tmp_path / "plain")
    mcp_paths = export_report(config, compare, tmp_path / "middle" / "mcp-outputs")
    server_paths = export_report(
        config,
        compare,
        tmp_path / "direct" / "mcp-outputs" / "config-audit",
    )

    assert Path(plain_paths.excel_path).is_relative_to(
        tmp_path / "plain" / "mcp-outputs" / "config-audit"
    )
    assert Path(mcp_paths.excel_path).is_relative_to(
        tmp_path / "middle" / "mcp-outputs" / "config-audit"
    )
    assert Path(server_paths.excel_path).is_relative_to(
        tmp_path / "direct" / "mcp-outputs" / "config-audit"
    )


def _config() -> NormalizedConfig:
    return NormalizedConfig(
        device_profile=DeviceProfile(
            device_name="fw-a",
            vendor="H3C",
            model="F1000",
            os_version="V7",
            site_name="测试站点",
            local_area_name="DMZ",
            standard_zone="dmz",
            role="border_firewall",
        ),
        zones=[Zone(name="trust"), Zone(name="untrust")],
        interfaces=[Interface(name="GigabitEthernet1/0/1", ip_addresses=["10.0.0.1/24"])],
        address_objects=[AddressObject(name="web_server", values=["10.0.0.10"])],
        service_objects=[ServiceObject(name="https", protocol="tcp", ports=["443"])],
        policy_rules=[
            PolicyRule(
                name="allow_any",
                action="permit",
                source_objects=["any"],
                destination_objects=["any"],
                services=["any"],
                logging=False,
            )
        ],
        management_access=[ManagementAccess(protocol="ssh", allowed_sources=["10.0.0.0/24"])],
    )
