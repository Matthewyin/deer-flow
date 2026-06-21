from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

from core.model import (
    AddressObject,
    CompareResult,
    Interface,
    LoggingConfig,
    ManagementAccess,
    NatRule,
    NormalizedConfig,
    PolicyRule,
    ReportPaths,
    Route,
    ServiceObject,
    UnparsedBlock,
    Zone,
)


EXCEL_FILENAME = "配置审查事实表.xlsx"
MARKDOWN_FILENAME = "配置审查报告.md"
VIRTUAL_OUTPUT_PREFIX = "/mnt/user-data/outputs/.mcp/config-audit"

MODULE_SHEETS = [
    "zones",
    "interfaces",
    "address_objects",
    "service_objects",
    "policy_rules",
    "nat_rules",
    "routes",
    "management_access",
    "logging",
    "unparsed_blocks",
]

MODULE_MODELS: dict[str, type[BaseModel]] = {
    "zones": Zone,
    "interfaces": Interface,
    "address_objects": AddressObject,
    "service_objects": ServiceObject,
    "policy_rules": PolicyRule,
    "nat_rules": NatRule,
    "routes": Route,
    "management_access": ManagementAccess,
    "logging": LoggingConfig,
    "unparsed_blocks": UnparsedBlock,
}

FINDING_HEADERS = [
    "finding_id",
    "severity",
    "module",
    "finding_type",
    "expected",
    "actual",
    "evidence",
    "llm_analysis",
    "recommendation",
]


def export_report(
    config: NormalizedConfig,
    compare: CompareResult,
    output_dir: str | Path,
    agent_analysis: str = "",
) -> ReportPaths:
    report_dir = _create_report_dir(output_dir)
    excel_path = report_dir / EXCEL_FILENAME
    markdown_path = report_dir / MARKDOWN_FILENAME

    _write_excel(config, compare, excel_path)
    _write_markdown(config, compare, markdown_path, agent_analysis)

    paths = _present_filepaths(report_dir.name)
    return ReportPaths(
        excel_path=str(excel_path),
        markdown_path=str(markdown_path),
        present_filepaths=paths,
    )


def _create_report_dir(output_dir: str | Path) -> Path:
    root = _resolve_output_root(output_dir)
    root.mkdir(parents=True, exist_ok=True)
    for _ in range(10):
        name = f"config-audit-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid4().hex[:8]}"
        path = root / name
        try:
            path.mkdir()
            return path
        except FileExistsError:
            continue
    raise FileExistsError("无法创建唯一的配置审查报告目录")


def _resolve_output_root(output_dir: str | Path) -> Path:
    root = Path(output_dir)
    if root.name == "config-audit" and root.parent.name == "mcp-outputs":
        return root
    if root.name == "mcp-outputs":
        return root / "config-audit"
    return root / "mcp-outputs" / "config-audit"


def _present_filepaths(artifact_id: str) -> list[str]:
    return [
        f"{VIRTUAL_OUTPUT_PREFIX}/{artifact_id}/{EXCEL_FILENAME}",
        f"{VIRTUAL_OUTPUT_PREFIX}/{artifact_id}/{MARKDOWN_FILENAME}",
    ]


def _write_excel(config: NormalizedConfig, compare: CompareResult, excel_path: Path) -> None:
    workbook = Workbook()
    device_sheet = workbook.active
    device_sheet.title = "device_profile"
    _write_device_profile(device_sheet, config)

    for sheet_name in MODULE_SHEETS:
        worksheet = workbook.create_sheet(sheet_name)
        _write_module_sheet(worksheet, sheet_name, getattr(config, sheet_name))

    findings_sheet = workbook.create_sheet("findings")
    _write_findings_sheet(findings_sheet, compare)

    workbook.save(excel_path)


def _write_device_profile(worksheet: Worksheet, config: NormalizedConfig) -> None:
    worksheet.append(["field", "value"])
    for field, value in config.device_profile.model_dump().items():
        worksheet.append([field, _cell_value(value)])


def _write_module_sheet(worksheet: Worksheet, sheet_name: str, items: list[BaseModel]) -> None:
    headers = _headers_for_items(sheet_name, items)
    worksheet.append(headers)
    if not items:
        worksheet.append(_placeholder_row(headers))
        return

    for item in items:
        data = item.model_dump()
        worksheet.append([_cell_value(data.get(header, "")) for header in headers])


def _write_findings_sheet(worksheet: Worksheet, compare: CompareResult) -> None:
    worksheet.append(FINDING_HEADERS)
    if not compare.findings:
        worksheet.append(["无发现", "info", "", "", "", "", "", "", ""])
        return

    for finding in compare.findings:
        data = finding.model_dump()
        worksheet.append([_cell_value(data.get(header, "")) for header in FINDING_HEADERS])


def _headers_for_items(sheet_name: str, items: list[BaseModel]) -> list[str]:
    if items:
        return list(items[0].__class__.model_fields)
    return list(MODULE_MODELS[sheet_name].model_fields)


def _placeholder_row(headers: list[str]) -> list[str]:
    if not headers:
        return ["无数据"]
    return ["无数据", *["" for _ in headers[1:]]]


def _cell_value(value: Any) -> str | int | float | bool | None:
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    if isinstance(value, dict):
        return "\n".join(f"{key}: {item}" for key, item in value.items())
    return value


def _write_markdown(
    config: NormalizedConfig,
    compare: CompareResult,
    markdown_path: Path,
    agent_analysis: str,
) -> None:
    markdown_path.write_text(
        "\n".join(
            [
                "# 配置审查报告",
                "",
                "## 结论摘要",
                _summary(compare),
                "",
                "## 设备画像",
                _device_profile_markdown(config),
                "",
                "## 模块统计",
                _module_stats_markdown(config),
                "",
                "## 关键缺失项",
                _missing_findings_markdown(compare),
                "",
                "## 异常与风险项",
                _risk_findings_markdown(compare),
                "",
                "## LLM 分析说明",
                _llm_analysis_markdown(compare, agent_analysis),
                "",
                "## 整改建议",
                _recommendations_markdown(compare),
                "",
            ]
        ),
        encoding="utf-8",
    )


def _summary(compare: CompareResult) -> str:
    if not compare.findings:
        return "未发现配置缺失或风险项。"
    severity_counts: dict[str, int] = {}
    for finding in compare.findings:
        severity_counts[finding.severity] = severity_counts.get(finding.severity, 0) + 1
    count_text = "，".join(
        f"{severity} {count} 项" for severity, count in sorted(severity_counts.items())
    )
    return f"本次审查发现 {len(compare.findings)} 项问题：{count_text}。"


def _device_profile_markdown(config: NormalizedConfig) -> str:
    lines = ["| 字段 | 值 |", "| --- | --- |"]
    for field, value in config.device_profile.model_dump().items():
        lines.append(f"| {field} | {_markdown_cell(value)} |")
    return "\n".join(lines)


def _module_stats_markdown(config: NormalizedConfig) -> str:
    lines = ["| 模块 | 数量 |", "| --- | ---: |"]
    for module in MODULE_SHEETS:
        lines.append(f"| {module} | {len(getattr(config, module))} |")
    return "\n".join(lines)


def _missing_findings_markdown(compare: CompareResult) -> str:
    findings = [
        finding
        for finding in compare.findings
        if finding.finding_type == "missing"
    ]
    return _finding_list(findings, "无关键缺失项。")


def _risk_findings_markdown(compare: CompareResult) -> str:
    findings = [
        finding
        for finding in compare.findings
        if finding.finding_type != "missing"
    ]
    return _finding_list(findings, "无异常与风险项。")


def _llm_analysis_markdown(compare: CompareResult, agent_analysis: str) -> str:
    lines: list[str] = []
    if agent_analysis:
        lines.append(agent_analysis)
    else:
        lines.append("未提供额外 LLM 分析。")

    finding_analysis = [
        f"- {finding.finding_id}：{finding.llm_analysis}"
        for finding in compare.findings
        if finding.llm_analysis
    ]
    if finding_analysis:
        lines.extend(["", *finding_analysis])
    return "\n".join(lines)


def _recommendations_markdown(compare: CompareResult) -> str:
    recommendations = [
        f"- {finding.finding_id}：{finding.recommendation}"
        for finding in compare.findings
        if finding.recommendation
    ]
    if not recommendations:
        return "暂无整改建议。"
    return "\n".join(recommendations)


def _finding_list(findings: list[Any], empty_text: str) -> str:
    if not findings:
        return empty_text
    return "\n".join(
        f"- {finding.finding_id} [{finding.severity}] {finding.actual}；证据：{finding.evidence}"
        for finding in findings
    )


def _markdown_cell(value: Any) -> str:
    text = str(_cell_value(value) or "")
    return text.replace("\n", "<br>").replace("|", "\\|")
